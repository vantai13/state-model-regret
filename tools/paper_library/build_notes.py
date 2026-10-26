"""Build the small current paper workbook, preserving author progress and notes.

Responsible: Doan Van Tai. Reads papers/data/catalog.json; no network/downloads.
The complete former library is frozen under papers/archive/.
"""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
DEST = ROOT / "papers"
DATA = DEST / "data"
LOCAL = Path.home() / "papers/state-model-regret/current"
OUTPUT = DEST / "paper_comparison.xlsx"

LIST_HEADERS = ["ID", "Tên bài", "Tác giả", "Năm", "Nơi đăng", "Hình thức công bố", "Loại paper",
                "Peer-reviewed?", "Vùng", "Ưu tiên", "Liên quan tới claim", "Mức tôi đã đọc",
                "PDF / Truy cập", "DOI / Link", "Vai trò hiện tại"]
NOTE_HEADERS = ["ID", "Tên bài", "Vấn đề", "Cách làm", "Kết quả chính + vị trí", "Giới hạn",
                "Khác với đề tài tôi", "Overlap E/M/R", "Mức AI đã đọc", "Dataset / Code",
                "Citations (nguồn, ngày)", "BibTeX key", "Việc tiếp theo", "Ghi chú do", "Updated at"]
AUTHOR_LEVELS = ["Chưa có xác nhận", "Chưa đọc", "Metadata", "Abstract", "Mục liên quan", "Full bài"]
AI_LEVELS = ["Thứ cấp", "Metadata", "Abstract", "Mục liên quan", "Full bài"]


def rows():
    data = json.loads((DATA / "catalog.json").read_text())["papers"]
    ids = [d["id"] for d in data]
    assert ids and len(ids) == len(set(ids)), "Paper IDs must be unique"
    return data


def style(sheet, widths, height=72):
    sheet.freeze_panes = "C2"
    thin = Side(style="thin", color="BFBFBF")
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="000000")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, color="000000", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    sheet.row_dimensions[1].height = 30
    for r in range(2, sheet.max_row + 1):
        sheet.row_dimensions[r].height = height
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_title_rows = "1:1"


def add_list(sheet, column, values):
    rule = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    rule.showErrorMessage = True
    sheet.add_data_validation(rule)
    rule.add(f"{column}2:{column}{sheet.max_row}")


def build(data):
    previous_reading, previous_notes = {}, {}
    if OUTPUT.exists():
        previous = load_workbook(OUTPUT)
        for row in previous["Danh_sach"].iter_rows(min_row=2, values_only=True):
            previous_reading[row[0]] = row[11]
        for row in previous["Ghi_chu_doc"].iter_rows(min_row=2, values_only=True):
            previous_notes[row[0]] = list(row)
        previous.close()
        # A catalog edit must not silently discard personal records on rebuild.
        assert (set(previous_notes) | set(previous_reading)) <= {d["id"] for d in data}, "Archive removed rows before rebuilding"

    book = Workbook()
    book.properties.creator = book.properties.lastModifiedBy = "Doan Van Tai"
    book.properties.title = "Paper — danh sách đang dùng"
    listing = book.active
    listing.title = "Danh_sach"
    listing.append(LIST_HEADERS)
    notes = book.create_sheet("Ghi_chu_doc")
    notes.append(NOTE_HEADERS)
    for d in data:
        listing.append([d["id"], d["title"], d["authors"], d["year"], d["venue"], d["type"], d["kind"],
                        d["peer"], d["area"], d["priority"], d["claims"],
                        previous_reading.get(d["id"], d["author_reading"]),
                        d["access"], d["url"].replace("https://doi.org/", "doi:"), d["role"]])
        notes.append(previous_notes.get(d["id"], d["note_row"]))
    style(listing, [9, 43, 24, 7, 25, 15, 19, 13, 26, 10, 42, 20, 23, 31, 46])
    style(notes, [9, 35, 35, 35, 44, 33, 35, 12, 17, 16, 25, 24, 35, 18, 15], 84)
    for r, d in enumerate(data, 2):
        listing.cell(r, 14).hyperlink = d["url"]
        listing.cell(r, 14).font = Font(name="Calibri", size=11, color="000000", underline="single")
        if d["local"]:
            listing.cell(r, 13).hyperlink = "PDF_local/" + d["local"]
            listing.cell(r, 13).font = Font(name="Calibri", size=11, color="000000", underline="single")
    add_list(listing, "L", AUTHOR_LEVELS)
    add_list(notes, "I", AI_LEVELS)
    add_list(notes, "N", ["AI (chưa kiểm)", "Tôi"])
    guide = book.create_sheet("Huong_dan")
    guide.append(["Mục", "Cách dùng"])
    for row in [
        ("Phạm vi", f"{len(data)} nguồn trong catalog hiện dùng; xem cột Vai trò hiện tại. Không phải danh sách đầy đủ về NDT hay bài tập phải đọc hết."),
        ("L2.6 Part A", "S11 hiện tại: freshness → decision cost; bản đọc arxivv3, journal identity TIT2026 đã kiểm. K3 ở notes/literature/freshness_decision_value.md; P13/Part B, tác giả L-self/Gate còn mở."),
        ("L2.5 Part2", "M00 → P04: herding theo tin cũ và fixed-M0 bias. Part2 đã kiểm kỹ thuật; tác giả L-self/Gate còn mở. P01 chỉ targeted screen, chưa cần đọc full."),
        ("L2.5 Part1", "P04: stale link-state, blocking và cost levels; Fig8 không phải clean2×2. Mức tác giả đọc cập nhật riêng, không tự nâng theo AI."),
        ("L2.4", "Nguồn lesson trước: NDT1 → P05 → P10. NMRG, P08, S07 tra bổ sung trong notes/literature/ndt_landscape.md."),
        ("R32 / R34", "Nguồn gần trục state: feedback control và queue với thông tin trễ; khác bài toán routing của repo."),
        ("R51 = P03", "DT² là cầu nối phương pháp digital twin, không phải paper networking chuyên biệt."),
        ("Mức đọc", "Bạn cập nhật Danh_sach → Mức tôi đã đọc; mức AI riêng ở Ghi_chu_doc. Giữ nguyên record cũ, không tự tăng mức đọc."),
        ("Ghi chú", "Ghi_chu_doc giữ nội dung theo ID; các nhắc lịch lesson trong ghi chú cũ là lịch sử. Mức AI đọc luôn tách khỏi mức bạn đọc."),
        ("Tìm paper", "Dùng notes/literature/search_protocol.md; tác giả duyệt trước khi tải hoặc tạo note mới."),
        ("Lưu tạm", "archive/paper_comparison_part_a.xlsx giữ nguyên 21 bài và toàn bộ ghi chú trước dọn. Không phải yêu cầu đọc hết."),
        ("Nguồn lý thuyết", "SOB14, B82, P11 chỉ tra tên khái niệm khi cần; note/proof Part A vẫn ở notes/literature/."),
        ("Link", "Bấm DOI / Link hoặc PDF / Truy cập. Tên bài không gắn link; PDF chỉ mở trên máy có thư viện local."),
        ("E/M/R / ưu tiên", "Ước lượng AI từ lượt cũ; không tự chứng nhận novelty hay đổi nhãn PC/Gate."),
    ]:
        guide.append(row)
    style(guide, [23, 110], 38)
    book.save(OUTPUT)
    book.close()


def verify(data):
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        assert not any(n.lower().endswith(".pdf") for n in archive.namelist())
    book = load_workbook(OUTPUT)
    assert book.sheetnames == ["Danh_sach", "Ghi_chu_doc", "Huong_dan"]
    ids = [d["id"] for d in data]
    for name in ["Danh_sach", "Ghi_chu_doc"]:
        sheet = book[name]
        assert sheet.max_row == len(ids) + 1
        assert [r[0] for r in sheet.iter_rows(min_row=2, values_only=True)] == ids
        assert all(sheet.cell(r, 2).hyperlink is None for r in range(2, sheet.max_row + 1))
        fills = {c.fill.fgColor.rgb for row in sheet for c in row if c.fill.fill_type == "solid"}
        assert fills <= {"00D9D9D9"}, fills
    for r, d in enumerate(data, 2):
        assert book["Danh_sach"].cell(r, 12).value in AUTHOR_LEVELS
        assert book["Ghi_chu_doc"].cell(r, 9).value in AI_LEVELS
        assert book["Danh_sach"].cell(r, 14).hyperlink.target == d["url"]
        if d["local"]:
            pdf = LOCAL / d["local"]
            assert pdf.read_bytes().startswith(b"%PDF-")
            assert hashlib.sha256(pdf.read_bytes()).hexdigest() == d["local_sha256"]
            assert (DEST / book["Danh_sach"].cell(r, 13).hyperlink.target).resolve() == pdf.resolve()
    result = dict(status="PASS", responsible="Doan Van Tai", workbook=OUTPUT.name,
                  papers=len(ids), ids=ids, sheets=book.sheetnames,
                  local_pdfs=sum(bool(d["local"]) for d in data), embedded_pdfs=0,
                  title_links=0, fill_colours=["D9D9D9"],
                  author_reading_unconfirmed=sum(book["Danh_sach"].cell(r, 12).value == "Chưa có xác nhận"
                                                for r in range(2, len(ids) + 2)),
                  workbook_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest())
    book.close()
    (DATA / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False), flush=True)
    return result


if __name__ == "__main__":
    records = rows()
    build(records)
    verify(records)

"""Build the concise Part A workbook and an optional local-only PDF library.

Responsible: Doan Van Tai. Summaries/comparisons are AI-assisted proposals.
PDFs are stored under ~/papers/, never copied into the Git repository.
"""
import argparse
import csv
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
import urllib.request
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
DEST = ROOT / "papers"
DATA = DEST / "data"
LOCAL = Path.home() / "papers" / "state-model-regret"
LEVELS = ["Chưa đọc", "Metadata", "Abstract", "Mục liên quan", "Full bài"]


def download(paper):
    target = LOCAL / paper["pdf_file"]
    target.parent.mkdir(parents=True, exist_ok=True)
    errors = []
    old = Path.home() / "papers" / "l2_part_a" / (paper["id"] + ".pdf")
    urls = paper["pdf_urls"]
    data, source = None, None
    if target.is_file():
        data, source = target.read_bytes(), "existing local copy"
    elif old.is_file():
        data, source = old.read_bytes(), urls[0]
    else:
        for url in urls:
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=20) as response:
                    candidate = response.read()
                if not candidate.startswith(b"%PDF-"):
                    raise ValueError("Response is not a PDF")
                data, source = candidate, url
                break
            except Exception as error:
                errors.append(type(error).__name__ + ": " + str(error))
    if data:
        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid cached PDF: " + paper["id"])
        # Validate a complete PDF and its identity before recording success.
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            trial = Path(tmp) / "paper.pdf"
            trial.write_bytes(data)
            info = subprocess.run(["pdfinfo", str(trial)], capture_output=True, text=True, check=True).stdout
            first = subprocess.run(["pdftotext", "-f", "1", "-l", "1", str(trial), "-"], capture_output=True, text=True, check=True).stdout
        tokens = set(re.findall(r"[a-z]+", paper["title"].lower()))
        observed = set(re.findall(r"[a-z]+", first.lower()))
        overlap = len(tokens & observed) / len(tokens)
        if overlap < .70:
            raise ValueError("PDF title needs manual review: " + paper["id"])
        pages = int(next(line.split(":", 1)[1] for line in info.splitlines() if line.startswith("Pages:")))
        if not target.exists():
            target.write_bytes(data)
        return dict(id=paper["id"], status="Đã tải", file=paper["pdf_file"],
                    pages=pages, bytes=len(data), sha256=hashlib.sha256(data).hexdigest(),
                    source=source, title_overlap=round(overlap, 3), errors=errors)
    return dict(id=paper["id"], status="Chưa tải được", file="", errors=errors,
                source=paper["source_url"])


def style_sheet(sheet, widths, table_name):
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=11, color="243746")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sheet.row_dimensions[1].height = 32
    for i in range(2, sheet.max_row + 1):
        sheet.row_dimensions[i].height = 88
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "1:1"


def build(papers, status):
    book = Workbook()
    book.properties.creator = "Doan Van Tai"
    book.properties.lastModifiedBy = "Doan Van Tai"
    book.properties.title = "Part A — Paper comparison"
    book.properties.description = "AI-assisted summaries; reading progress does not attest author reading."
    book.properties.created = datetime(2026, 10, 24)
    main = book.active
    main.title = "So_sanh_Part_A"
    main.append(["ID", "Ưu tiên", "Topic", "Năm", "Tác giả", "Tên bài", "Claim chính", "So với nghiên cứu của tôi", "Giới hạn cần nhớ", "Nguồn", "AI đã đọc"])
    for paper in papers:
        main.append([paper[k] for k in ("id", "priority", "topic", "year", "authors_short", "title", "claim", "compare", "limit", "venue", "ai_reading")])
        r = main.max_row
        main.cell(r, 6).hyperlink = paper["source_url"]
        main.cell(r, 6).comment = Comment("Nguồn: " + paper["source_url"] + "\nVị trí: " + paper["read_location"] + "\nSo sánh là AI đề xuất.", "Doan Van Tai")
        main.cell(r, 10).hyperlink = paper["source_url"]
        main.cell(r, 5).comment = Comment(paper["authors"], "Doan Van Tai")
    style_sheet(main, [11, 13, 23, 9, 29, 46, 43, 46, 42, 28, 18], "PartAPapers")
    progress = book.create_sheet("Tien_do_doc")
    progress.append(["ID", "AI đã đọc", "Phần đã đọc / căn cứ", "Bạn đã đọc", "Trạng thái nguồn", "PDF cục bộ", "Đọc tiếp", "DOI / nguồn chính thức", "Ngày kiểm"])
    for paper in papers:
        pdf = status[paper["id"]]
        label = pdf["status"] + (f" ({pdf['pages']} trang)" if pdf.get("pages") else "")
        progress.append([paper["id"], paper["ai_reading"], paper["read_location"], paper["author_reading"], paper["source_type"], label, paper["next_read"], paper["source_url"], paper["verified_date"]])
        r = progress.max_row
        progress.cell(r, 8).hyperlink = paper["source_url"]
        if pdf["status"] == "Đã tải":
            progress.cell(r, 6).hyperlink = "PDF_local/" + pdf["file"]
        else:
            progress.cell(r, 6).hyperlink = paper["pdf_urls"][0]
            progress.cell(r, 6).comment = Comment("; ".join(pdf["errors"]), "Doan Van Tai")
    style_sheet(progress, [11, 18, 48, 23, 29, 25, 48, 48, 14], "ReadingProgress")
    colors = {"Mục liên quan": "E2F0D9", "Abstract": "FFF2CC", "Metadata": "E7E6E6"}
    for r, paper in enumerate(papers, 2):
        fill = PatternFill("solid", fgColor=colors.get(paper["ai_reading"], "FFFFFF"))
        main.cell(r, 11).fill = fill
        progress.cell(r, 2).fill = fill
        progress.cell(r, 4).fill = PatternFill("solid", fgColor="F2F2F2")
    validation = DataValidation(type="list", formula1='"Chưa có xác nhận,Chưa đọc,Metadata,Abstract,Mục liên quan,Full bài"')
    validation.errorTitle = "Chọn mức đọc"
    validation.error = "Chọn một mức trong danh sách."
    validation.showErrorMessage = True
    progress.add_data_validation(validation)
    validation.add(f"D2:D{len(papers)+1}")
    guide = book.create_sheet("Huong_dan")
    guide.append(["Mục", "Cách hiểu"])
    for row in [
        ("Phụ trách", "Doan Van Tai; cập nhật 24/10/2026."),
        ("Phạm vi", "Chỉ 6 bài Part A. Ba bài cốt lõi: SOB14, B82, P11."),
        ("Đọc bảng", "Mở So_sanh_Part_A; bấm tên bài để mở nguồn. Dùng bộ lọc ở header."),
        ("Mức đọc", "Metadata < Abstract < Mục liên quan < Full bài. Có PDF không có nghĩa đã đọc."),
        ("Ai đã đọc", "Cột AI và cột Bạn tách riêng. Bạn tự cập nhật cột Bạn đã đọc khi đọc thật."),
        ("Nguồn", "Tạp chí/hội nghị chính thức khác working paper. Không gán peer review cho MAB11 WP."),
        ("PDF", "Bấm PDF để mở bản cục bộ hoặc URL. PDF_local chỉ có trên máy tạo thư viện; không push PDF."),
        ("Nội dung", "Claim theo mức đọc đã ghi; so sánh/ưu tiên là AI đề xuất, không xác nhận novelty."),
        ("Cập nhật", "Sửa Excel trực tiếp. Builder tạo lại từ data/catalog.json; lưu thay đổi trước khi chạy."),
    ]:
        guide.append(row)
    style_sheet(guide, [22, 115], "Guide")
    for r in range(2, guide.max_row+1):
        guide.row_dimensions[r].height = 35
    # Restore explicit hyperlink appearance after applying common typography.
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.hyperlink:
                    cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
        sheet.sheet_properties.tabColor = "17365D" if sheet == main else "548235"
    output = DEST / "paper_comparison_part_a.xlsx"
    book.save(output)
    return output


def verify(output, papers, status):
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert not any(name.endswith(".pdf") for name in archive.namelist())
    book = load_workbook(output)
    assert book.sheetnames == ["So_sanh_Part_A", "Tien_do_doc", "Huong_dan"]
    ids = [p["id"] for p in papers]
    assert len(ids) == len(set(ids)) == 6
    assert all(p["part"] == "Part_A" for p in papers)
    for sheet in list(book)[:2]:
        assert sheet.max_row == 7
        assert [sheet.cell(r, 1).value for r in range(2, 8)] == ids
        assert sheet.freeze_panes == "C2" and len(sheet.tables) == 1
    assert all(book["Tien_do_doc"].cell(r, 4).value == "Chưa có xác nhận" for r in range(2, 8))
    assert not any(p["ai_reading"] == "Full bài" for p in papers)
    for r, paper in enumerate(papers, 2):
        assert book["So_sanh_Part_A"].cell(r, 6).value == paper["title"]
        assert book["So_sanh_Part_A"].cell(r, 6).hyperlink.target == paper["source_url"]
        assert book["So_sanh_Part_A"].cell(r, 4).value == paper["year"]
        assert len(paper["claim"]) <= 140 and len(paper["compare"]) <= 150
        row = status[paper["id"]]
        if row["status"] == "Đã tải":
            assert hashlib.sha256((LOCAL / row["file"]).read_bytes()).hexdigest() == row["sha256"]
    with (DEST / "part_a_quick_view.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert [r["ID"] for r in csv.DictReader(stream)] == ids
    counts = {level: sum(p["ai_reading"] == level for p in papers) for level in LEVELS}
    result = dict(status="PASS", papers=len(ids), sheets=book.sheetnames,
                  reading_counts=counts, pdf_downloaded=sum(s["status"] == "Đã tải" for s in status.values()),
                  no_part_b=True, author_reading_unconfirmed=6,
                  csv_rows_match=True, xlsx_archive_valid=True, embedded_pdfs=0,
                  workbook_sha256=hashlib.sha256(output.read_bytes()).hexdigest())
    (DATA / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False), flush=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download", action="store_true")
    args = parser.parse_args()
    catalog = json.loads((DATA / "catalog.json").read_text())
    papers = catalog["papers"]
    assert len(papers) == 6 and all(p["part"] == "Part_A" for p in papers)
    if args.download:
        LOCAL.mkdir(parents=True, exist_ok=True)
        with ThreadPoolExecutor(max_workers=4) as pool:
            records = list(pool.map(download, papers))
        (DATA / "pdf_status.json").write_text(json.dumps(records, ensure_ascii=False, indent=2) + "\n")
        alias = DEST / "PDF_local"
        if not alias.exists() and not alias.is_symlink():
            alias.symlink_to(LOCAL, target_is_directory=True)
        readme = LOCAL / "Part_A" / "README.md"
        readme.write_text("# Part A — State × Model Regret\n\nDoan Van Tai. PDF cục bộ; không push GitHub.\n\n" + "\n".join(
            f"- {p['id']}: {p['title']} — {s['status']}. {p['source_url']}"
            for p, s in zip(papers, records)) + "\n")
    records = json.loads((DATA / "pdf_status.json").read_text())
    status = {r["id"]: r for r in records}
    output = build(papers, status)
    with (DEST / "part_a_quick_view.csv").open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["ID", "Ưu tiên", "Topic", "Năm", "Tác giả", "Claim chính", "So với nghiên cứu của tôi", "AI đã đọc", "Nguồn"])
        for paper in papers:
            writer.writerow([paper[k] for k in ("id", "priority", "topic", "year", "authors_short", "claim", "compare", "ai_reading", "source_url")])
    verify(output, papers, status)
    print("Workbook:", output, flush=True)


if __name__ == "__main__":
    main()

"""Audit lượt đối chiếu bản chính thức S11/R01/R34, gói tự đọc và ví dụ T03 (26/10/2026).

Người phụ trách: Doan Van Tai. Kiểm hash/danh tính, trùng TR của R01, số mục S11 journal, ví dụ
T03 tính lại đúng và không hòa, catalog chỉ đổi 3 ID, ô của tác giả (L-self, nhãn) vẫn trống,
chỉ append/insert, link sống, không PDF trong Git.
"""
from __future__ import annotations

import difflib
import hashlib
import json
import re
import subprocess
import sys
from fractions import Fraction as F
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results/lessons/l2_9/journal_check"
BASE = "c3e7ea7c741931a0d3452c10093df7fb12b91d9e"
sys.path.insert(0, str(ROOT / "experiments/lessons"))
from t03_worked_example import run  # noqa: E402

APPEND_ONLY = ["notes/literature/prior_boundary.md", "notes/literature/papers/S11.md", "notes/literature/papers/R01.md",
               "notes/literature/search_log.md", "notes/00_evidence_inventory.md", "notes/00_exposure_register.md",
               "AGENTS.md", "notes/literature/author_reading_record.md"]


def git_show(p: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{BASE}:{p}"], cwd=ROOT)


def text(pdf: Path) -> str:
    return re.sub(r"\s+", " ", subprocess.check_output(["pdftotext", str(pdf), "-"], text=True))


def main() -> None:
    acc = json.loads((OUT / "source_access.json").read_text(encoding="utf-8"))
    I = Path(acc["intake_dir"]); assert not I.resolve().is_relative_to(ROOT)
    for s in acc["sources"]:
        p = I / s["file"]
        assert hashlib.sha256(p.read_bytes()).hexdigest() == s["sha256"]
        assert int(re.search(r"Pages:\s+(\d+)", subprocess.check_output(["pdfinfo", str(p)], text=True))[1]) == s["pages"]
    s11 = text(I / "S11_Li2026_TIT72-4_2237.pdf")
    for q in ["VOL. 72, NO. 4, APRIL 2026", "Theorem 9 (Sensitivity Analysis)", "Corollary 2 (Sampling frequency threshold)",
              "found in Appendix O", "(10a)", "S YSTEM M ODEL AND P ROBLEM F ORMULATION"]:
        assert q in s11, q
    tr = text(Path.home() / "papers/state-model-regret/current/R01_Chen2006_TR.pdf")
    ec = text(I / "R01_Chen2006TR_eCommons_cover.pdf")
    ec = ec[ec.find("Re-Solving Stochastic Programming Models for Airline Revenue Management ∗"):]
    assert difflib.SequenceMatcher(None, tr[:20000], ec[:20000], autojunk=False).ratio() > 0.999 and "June 22, 2006" in ec
    r34 = text(I / "R34_KuriKumar1992_CDC_2pages.pdf")
    assert "Kuri" in r34 and "the JSEQ policy is optimal" in r34 and "JSEQ is no longer optimal" in r34

    r = run(1, 1, 4, 1, 4, F(1, 3), F(2, 3))
    assert (r["L11"], r["L10"], r["L01"], r["L00"]) == (F(8, 9), F(10, 9), F(1), F(10, 9)) and not r["any_tie"]
    assert r["L01"] - r["L11"] == r["evppi"] == F(1, 9) and r["L10"] - r["L11"] == r["Evss"] == F(2, 9)
    v = run(1, 1, 4, 1, 4, F(2, 3), F(2, 3))
    assert (v["L11"], v["L01"], v["evppi"], v["Evss"]) == (F(7, 9), F(8, 9), F(1, 9), F(1, 9))
    ex = json.loads((ROOT / "results/lessons/l2_9/t03_example/example.json").read_text(encoding="utf-8"))
    assert ex["EVPPI"] == "1/9" and ex["E_VSS"] == "2/9" and ex["Gamma"] == "1/9" and ex["exact_ties"] is False

    cat = json.loads((ROOT / "papers/data/catalog.json").read_text(encoding="utf-8"))["papers"]
    old = json.loads((OUT / "catalog_before.json").read_text(encoding="utf-8"))["papers"]
    assert (OUT / "catalog_before.json").read_bytes() == git_show("papers/data/catalog.json")
    assert sorted(n["id"] for n, o in zip(cat, old) if n != o) == ["R01", "R34", "S11"] and len(cat) == 23
    assert all(p["author_reading"] == "Chưa có xác nhận" for p in cat)
    ob = load_workbook(OUT / "paper_comparison_before.xlsx"); nb = load_workbook(ROOT / "papers/paper_comparison.xlsx")
    rows = lambda b: {x[0]: list(x) for x in b["Ghi_chu_doc"].iter_rows(min_row=2, values_only=True)}
    assert sorted(k for k, v in rows(ob).items() if rows(nb)[k] != v) == ["R01", "R34", "S11"]
    assert {x[11] for x in nb["Danh_sach"].iter_rows(min_row=2, values_only=True)} == {"Chưa có xác nhận"}

    rec = (ROOT / "notes/literature/author_reading_record.md").read_text(encoding="utf-8")
    trows = [x for x in rec.splitlines() if re.match(r"^\| [1-7] \|", x)]
    assert len(trows) == 7 and all(all(c.strip() == "" for c in x.strip().strip("|").split("|")[5:]) for x in trows)
    assert rec.count("L-2nd, author-attested") == 1
    pb = (ROOT / "notes/literature/prior_boundary.md").read_text(encoding="utf-8")
    assert pb.split("## L2.9 Part 1 — hồ sơ 8 cột", 1)[1].split("### Trạng thái Gate L", 1)[0].count("⟨tác giả chấm⟩") == 7
    for f in APPEND_ONLY:
        assert (ROOT / f).read_bytes().startswith(git_show(f)), f
    a = git_show("README.md").decode().splitlines(); b = (ROOT / "README.md").read_text(encoding="utf-8").splitlines()
    assert all(t in ("equal", "insert") for t, *_ in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes())

    assert not subprocess.check_output(["git", "ls-files", "*.pdf", "*.PDF"], cwd=ROOT).strip()
    links = 0
    for doc in [ROOT / "README.md", ROOT / "AGENTS.md", *OUT.glob("*.md"), *(ROOT / "notes").rglob("*.md")]:
        for link in re.findall(r"\]\(([^)]+)\)", doc.read_text(encoding="utf-8")):
            if "://" in link or link.startswith(("#", "mailto:")):
                continue
            t = unquote(link.split("#", 1)[0].strip("<>"))
            if (doc.parent / t).resolve() == (OUT / "verification.json").resolve():
                links += 1; continue
            assert (doc.parent / t).exists(), (str(doc), link); links += 1
    res = dict(artifact_status="PASS", responsible="Doan Van Tai", date="2026-10-26", base_commit=BASE,
               S11_journal_numbering_matches=True, R01_file_is_TR2006=True, R34_file_is_CDC1992=True,
               still_needed=acc["still_needed"], t03_example={"EVPPI": "1/9", "E_VSS": "2/9", "Gamma": "1/9", "ties": False},
               catalog_rows_changed=["R01", "R34", "S11"], author_levels_unconfirmed=23, author_L_self_records=0,
               author_label_cells_blank=7, provisional_labels_changed=0, local_links_checked=links, pdfs_in_git=0, gate_L="HOLD")
    (OUT / "verification.json").write_text(json.dumps(res, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(res, ensure_ascii=False))


if __name__ == "__main__":
    main()

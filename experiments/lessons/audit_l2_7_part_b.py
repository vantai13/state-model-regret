"""L2.7B audit: enumerate actions, replay metadata and preserve old artifacts.

Responsible: Doan Van Tai. Does not import the generator or run a plant.
"""
import csv
from fractions import Fraction as F
import hashlib
from itertools import product
import json
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"results/lessons/l2_7/part_b"


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def rows(name):
    with (OUT/name).open() as f:return list(csv.DictReader(f))


def main():
    result=json.loads((OUT/"toy_results.json").read_text())
    config=json.loads((OUT/"config.json").read_text())
    assert result["status"]=="PASS" and result["check_count"]==len(result["checks"])==27
    assert result["source_sha256"]==sha(ROOT/"experiments/lessons/l2_7_equivalence.py")
    for p,h in result["inputs"].items():assert sha(ROOT/p)==h,p
    assert config["selected_K_rule"] is None and config["MC_samples"]==0 and config["heldout_metadata_only"]
    sweep=rows("gap_bound_sweep.csv");assert len(sweep)==625;seen=set()
    for r in sweep:
        g=[F(r["true_gap1"]),F(r["true_gap2"])];h=[F(r["model_gap1"]),F(r["model_gap2"])];seen.add(tuple(g+h))
        truth=[[F(10),F(10)+x] for x in g];model=[[F(10),F(10)+x] for x in h]
        choose=lambda v:min(range(2),key=lambda a:v[a])
        a0=choose([sum(v[a] for v in model)/2 for a in range(2)])
        a1=choose([sum(v[a] for v in truth)/2 for a in range(2)])
        gm=sum(v[a0]-v[a1] for v in truth)/2
        selected=[choose(v) for v in model];best=[choose(v) for v in truth]
        known=sum(v[a]-v[b] for v,a,b in zip(truth,selected,best))/2
        assert F(r["GM"])==gm and F(r["GMgivenS1"])==known
        assert F(r["Gamma"])==known-gm
        assert F(r["S0_error"])==abs(sum(h)/2-sum(g)/2)
        assert F(r["S1_error"])==sum(abs(y-x) for x,y in zip(g,h))/2
        assert gm<=F(r["S0_error"]) and known<=F(r["S1_error"])
        cref=sum(min(v) for v in truth)/2
        assert F(r["Cref"])==cref and F(r["GM_relative"])==gm/cref
        assert F(r["GMgivenS1_relative"])==known/cref
    assert seen==set(product([F(-1,2),F(-1,4),F(0),F(1,4),F(1,2)],repeat=4))
    numeric=rows("numeric_intervals.csv");assert len(numeric)==7
    numeric_byid={r["case"]:r for r in numeric}
    for r in numeric:
        assert F(r["interval_lower"])==F(r["G"])-F(r["error"])
        assert F(r["interval_upper"])==F(r["G"])+F(r["error"])
    assert numeric_byid["closed_boundary"]["status"]=="EQUIVALENT_CERTIFIED"
    assert numeric_byid["lesson_straddle"]["status"]==numeric_byid["uncertain_reference"]["status"]=="NUMERICAL_INCONCLUSIVE"
    ci=rows("CI_examples.csv");assert len(ci)==4
    assert [r["case"] for r in ci if r["status"]=="EQUIVALENT_TOST_CRITERION"]==["pass"]
    assert F(result["examples"]["small_margin_exercise"]["GMgivenS1_relative"])==F(1,1999)
    assert result["V1"]["GM_relative"]=="0" and result["V1"]["GMgivenS1_relative"]=="1/5"

    coords=json.loads((OUT/"grid_coordinates.json").read_text());assert len(coords)==58
    keys={"id","role","kappa","alphaA","alphaB","eta","h","label_1","label_2"}
    assert all(set(c)==keys for c in coords)
    original=json.loads((ROOT/"notes/design/grid_v1_2.json").read_text())
    projected=[{k:r.get(k) for k in keys} for group in ["main","heldout","controls","probes"] for r in original[group]]
    assert coords==projected
    eligible={r["id"] for r in coords if r["role"] in {"main","heldout"} and r["label_1"]==r["label_2"]=="candidate"}
    wide={r["id"] for r in coords if r["id"] in eligible and F(r["kappa"])>=2 and
          r["alphaA"]==r["alphaB"]=="9/10" and F(r["h"])==8}
    assert wide=={"M19","M20","M21","M22","R03","R04"}
    membership=rows("K_membership.csv");assert len(membership)==174
    for rule,expected in [("listed_M19_M22",wide-{"R03","R04"}),
                          ("coordinate_kappa3",wide-{"R03","R04"}),
                          ("coordinate_kappa_ge2",wide)]:
        assert {r["id"] for r in membership if r["rule"]==rule and r["included"]=="True"}==expected

    protected=json.loads((OUT/"protected_before.json").read_text());assert len(protected["files"])==1163
    for p,h in protected["files"].items():assert sha(ROOT/p)==h,p
    for p in ["notes/literature/prior_boundary.md","notes/literature/already_implied.md",
              "notes/literature/matrix.md","notes/literature/search_log.md",
              "notes/00_exposure_register.md","notes/00_evidence_inventory.md"]:
        before=subprocess.check_output(["git","show",protected["baseline_commit"]+":"+p])
        assert (ROOT/p).read_bytes().startswith(before),p
    backup=Path.home()/"papers/state-model-regret/archive/2026-10-26_before_l2_7_part_b"
    old=load_workbook(backup/"paper_comparison.xlsx");new=load_workbook(ROOT/"papers/paper_comparison.xlsx")
    notes=lambda b:{r[0]:list(r) for r in b["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    progress=lambda b:{r[0]:r[11] for r in b["Danh_sach"].iter_rows(min_row=2,values_only=True)}
    before,after=notes(old),notes(new);assert len(before)==12 and set(after)-set(before)=={"P12"}
    assert all(after[k]==v for k,v in before.items())
    a,b=progress(old),progress(new);assert all(b[k]==v for k,v in a.items())
    assert new.sheetnames==["Danh_sach","Ghi_chu_doc","Huong_dan"]
    old.close();new.close()
    oldcat=json.loads((backup/"catalog.json").read_text())["papers"]
    cat=json.loads((ROOT/"papers/data/catalog.json").read_text())["papers"]
    assert cat[1:]==oldcat and cat[0]["id"]=="P12" and len({r["id"] for r in cat})==13
    access=json.loads((OUT/"source_access.json").read_text());p12=access["P12"]
    pdf=Path(p12["pdf_path"]);assert pdf.read_bytes().startswith(b"%PDF-") and sha(pdf)==p12["sha256"]
    assert not pdf.resolve().is_relative_to(ROOT) and p12["AI_level"]=="T"
    for im in p12["rendered_pages"]:assert sha(Path(im["path"]))==im["sha256"]
    r43=access["R43"];assert r43["AI_level_this_session"]=="A" and not r43["pdf_downloaded_this_session"]
    assert sha(Path(r43["archived_note"]))==r43["archived_note_sha256"]
    search=json.loads((OUT/"keyword_screen.json").read_text())
    assert search["screened_result_titles"]==len(search["hits"])==12 and search["kept_unique"]==0
    sums=(ROOT/"results/lessons/part_a/SHA256SUMS").read_text().splitlines();assert len(sums)==52
    for line in sums:
        h,p=line.split("  ",1);assert sha(ROOT/p)==h,p
    docs=[ROOT/"README.md",ROOT/"AGENTS.md",ROOT/"papers/README.md",OUT/"report.md",
          ROOT/"notes/lessons/L3.3_equivalence_queue.md",ROOT/"notes/00_exposure_register.md",
          ROOT/"notes/00_evidence_inventory.md",*list((ROOT/"notes/literature").rglob("*.md"))]
    links=0
    for doc in docs:
        for link in re.findall(r"\]\(([^)]+)\)",doc.read_text()):
            if "://" in link or link.startswith(("#","mailto:")):continue
            target=unquote(link.split("#",1)[0].strip("<>"))
            assert (doc.parent/target).exists(),(doc,link)
            links+=1
    assert not subprocess.check_output(["git","ls-files","*.pdf"]).strip()
    out=dict(artifact_status="PASS",part_b_technical_status="COMPLETE",
             whole_L2_7_status="OPEN_P02_FULL_TEXT_AND_HUMAN_GATE",responsible="Doan Van Tai",
             generator_checks=27,independent_two_state_replays=625,numeric_examples=7,CI_inputs=4,
             coordinate_rows=58,membership_rows=174,selected_K_rule=None,heldout_outcomes_used=False,
             protected_files_unchanged=1163,old_author_progress_preserved=12,old_notes_preserved=12,
             catalog_unique_papers=13,primary_P12_level="T",R43_current_level="A; archived AI-T provenance retained",
             legacy_manifest_files=52,relative_links_checked=links,MC_samples=0,grid_run=False,
             contract_changed=False,null_selected=False,human_gate="OPEN")
    (OUT/"verification.json").write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(out,ensure_ascii=False))


if __name__=="__main__":main()

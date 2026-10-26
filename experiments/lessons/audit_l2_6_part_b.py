"""Reconstruct L2.6 Part B artifacts and verify preservation.

Responsible: Doan Van Tai. Uses an arithmetic-series endpoint sum instead of
the generator's queue recurrence. Only writes part_b/verification.json.
"""
import csv
from fractions import Fraction as F
import hashlib
import json
from math import ceil
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"results/lessons/l2_6/part_b"


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def rows(name):
    with (OUT/name).open() as f:return list(csv.DictReader(f))


def main():
    result=json.loads((OUT/"toy_results.json").read_text())
    config=json.loads((OUT/"config.json").read_text())
    assert result["source_sha256"]==sha(ROOT/"experiments/lessons/l2_6_horizon.py")
    fluid=rows("fluid.csv");assert len(fluid)==8
    for row in fluid:
        q,r,h=map(F,(row["q0"],row["r"],row["h"]))
        end=max(F(0),q-r*h);full=q*q/(2*r)
        assert F(row["integral"])==full-end*end/(2*r)
        assert F(row["average"])==F(row["integral"])/h
        assert F(row["post_h_continuous_area"])==end*end/(2*r)
    endpoint=rows("endpoint_sum.csv");assert len(endpoint)==24
    for row in endpoint:
        q,r,h,dt=map(F,(row["q0"],row["r"],row["h"],row["dt"]))
        n=max(0,min(int(h/dt),ceil(q/(r*dt))-1))
        exact=dt*(n*q-r*dt*F(n*(n+1),2))
        assert F(row["endpoint_sum"])==exact
        assert F(row["qH"])==max(F(0),q-r*h)
        assert F(row["last_endpoint_charge"])==dt*F(row["qH"])
    gap_rows=rows("gap_horizon.csv");assert len(gap_rows)==6
    for row in gap_rows:
        gaps=[F(row["gap1"]),F(row["gap2"])];avg=sum(gaps)/2
        l01=10+min(F(0),avg)
        l11=10+sum((min(F(0),d) for d in gaps),F(0))/2
        assert F(row["L01"])==l01 and F(row["L11"])==l11
        assert F(row["GSgM1"])==l01-l11==min(F(row["Epositive"]),F(row["Enegative"]))
    for row in rows("continuous_optimizer.csv"):
        t=F(row["T"]);opt=F(row["mu_optimal"]);steady=F(row["mu_steady"])
        cost=lambda m:(m-5)**2+2*m/t
        assert opt==5-1/t
        assert F(row["average_loss"])==cost(steady)-cost(opt)==1/t**2
        assert F(row["cumulative_loss"])==1/t
    for row in rows("binary_step.csv"):
        gap=F(row["predicted_gap"])
        assert row["action"]==("Sw" if gap<0 else "St")
        assert F(row["loss"])==(F(0) if gap<0 else F(1,2))
    uniforms=rows("uniform_bias.csv");assert len(uniforms)==22
    for row in uniforms:
        a,b=F(row["halfwidth"]),F(row["bias"])
        assert F(row["loss"])==min(abs(b),a)**2/(4*a)
    finite=rows("finite_support_2x2.csv");assert len(finite)==22
    for row in finite:
        values=[F(v) for v in config["decision_examples"]["finite_supports"][row["case"]]]
        bias=F(row["bias"]);n=len(values);loss={}
        for cell in ("00","01","10","11"):
            total=F(0)
            for d in values:
                predicted=d if cell[0]=="1" else sum(values)/n
                if cell[1]=="0":predicted+=bias
                total+=5+(d if predicted<0 else 0)
            loss[cell]=total/n
            assert F(row["L"+cell])==loss[cell]
        assert F(row["GM"])==loss["00"]-loss["01"]
        assert F(row["GS"])==loss["00"]-loss["10"]
        assert F(row["GMgS1"])==loss["10"]-loss["11"]
        assert F(row["Gam"])==loss["10"]+loss["01"]-loss["00"]-loss["11"]
    moments=rows("moment_heuristic.csv");assert len(moments)==6
    for row in moments:
        second,mean,var=map(F,(row["second_moment"],row["mean"],row["variance"]))
        eta,r=F(row["eta"]),F(row["r_pre"])
        assert second==mean*mean+var
        assert F(row["B_only_limit"])==second/(2*r)*eta/(1-eta)
        assert F(row["mean_B_after_Switch"])<F(row["capacity"])
        support={"q=10":[F(10)],"q=20":[F(20)],"q=0_or20":[F(0),F(20)]}[row["distribution"]]
        h=F(config["moment"]["h_for_comparison"]);post=r*(1-eta)
        integral=lambda q,drift:(q*q-max(F(0),q-drift*h)**2)/(2*drift)
        finite_h=sum((integral(q,post)-integral(q,r) for q in support),F(0))/len(support)
        assert F(row["B_only_h8"])==finite_h

    protected=json.loads((OUT/"protected_before.json").read_text())
    assert all(sha(ROOT/p)==h for p,h in protected["files"].items())
    prefixes=["notes/literature/prior_boundary.md","notes/01_claim_boundary_v1_2.md",
              "notes/02_decision_log.md","notes/literature/h1_candidates.md"]
    for path in prefixes:
        before=subprocess.check_output(["git","show",protected["baseline_commit"]+":"+path])
        assert (ROOT/path).read_bytes().startswith(before),path
    backup=Path.home()/"papers/state-model-regret/archive/2026-10-26_before_l2_6_part_b"
    old=load_workbook(backup/"paper_comparison.xlsx");new=load_workbook(ROOT/"papers/paper_comparison.xlsx")
    records=lambda w:{r[0]:list(r) for r in w["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    before,after=records(old),records(new)
    assert len(before)==9 and set(after)-set(before)=={"P13"}
    assert all(after[k]==v for k,v in before.items())
    reading=lambda w:{r[0]:r[11] for r in w["Danh_sach"].iter_rows(min_row=2,values_only=True)}
    br,ar=reading(old),reading(new);assert all(ar[k]==v for k,v in br.items())
    old.close();new.close()
    old_catalog=json.loads((backup/"catalog.json").read_text())["papers"]
    catalog=json.loads((ROOT/"papers/data/catalog.json").read_text())["papers"]
    assert catalog[1:]==old_catalog and catalog[0]["id"]=="P13"
    access=json.loads((OUT/"source_access.json").read_text())["P13"]
    pdf=Path(access["path"])
    assert pdf.read_bytes().startswith(b"%PDF-") and sha(pdf)==access["sha256"]
    assert not pdf.resolve().is_relative_to(ROOT)
    sums=(ROOT/"results/lessons/part_a/SHA256SUMS").read_text().splitlines();assert len(sums)==52
    for line in sums:
        h,p=line.split("  ",1);assert sha(ROOT/p)==h,p
    docs=[ROOT/"README.md",ROOT/"AGENTS.md",ROOT/"papers/README.md",OUT/"report.md",
          ROOT/"notes/01_claim_boundary_v1_2.md",ROOT/"notes/02_decision_log.md",
          *list((ROOT/"notes/literature").rglob("*.md"))]
    links=0
    for path in docs:
        for link in re.findall(r"\]\(([^)]+)\)",path.read_text()):
            if "://" in link or link.startswith(("#","mailto:")):continue
            rel=unquote(link.split("#")[0].split(' "')[0])
            if not rel:continue
            target=(path.parent/rel).resolve()
            if target!=OUT/"verification.json":assert target.exists(),(path,rel)
            links+=1
    csv_files=sorted(OUT.glob("*.csv"));csv_count=sum(len(rows(p.name)) for p in csv_files)
    assert csv_count==102 and all(b"\r" not in p.read_bytes() for p in csv_files)
    verification=dict(status="PASS",responsible="Doan Van Tai",generator_checks=result["check_count"],
                      csv_files=len(csv_files),csv_rows=csv_count,fluid_rows=8,endpoint_rows=24,finite_tables=22,
                      previous_workbook_records_preserved=9,catalog_rows=10,
                      protected_files_unchanged=len(protected["files"]),old_document_prefixes_preserved=prefixes,
                      part_a_checksums_verified=52,local_links_checked=links,PDF_outside_git=True,
                      grid_outcomes_computed=0,P13_simulation_replication=False,formal_author_gate="pending")
    (OUT/"verification.json").write_text(json.dumps(verification,ensure_ascii=False,indent=2)+"\n")
    assert (OUT/"verification.json").exists()
    print(json.dumps(verification,ensure_ascii=False))


if __name__=="__main__":main()

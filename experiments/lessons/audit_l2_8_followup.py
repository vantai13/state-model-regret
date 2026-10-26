"""Audit artifact của L2.8 follow-up (26/10/2026).

Người phụ trách: Doan Van Tai. AI viết theo yêu cầu tác giả. Kiểm:
raw/ledger/CSV sàng khớp nhau theo rank; P02 PDF/hash/record; workbook và catalog
chỉ đổi đúng P02 so với snapshot trước follow-up; các note lịch sử chỉ được append
(hoặc thay đúng dòng P02 đã khai báo); artifact Part1/Part2 và contract không đổi;
link nội bộ còn sống; không PDF trong Git. Không chứng nhận search coverage,
novelty, L-self hay Gate.
"""

from __future__ import annotations

import csv
import difflib
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results" / "lessons" / "l2_8" / "followup"
BASE = "7f590236791f73dff8d8ff9f4db3df65eca1810f"
sys.path.insert(0, str(Path(__file__).resolve().parent))
from l2_8_followup_search import parse_records  # noqa: E402

DECISIONS = {"candidate_T1", "candidate_T2", "context_T3", "exclude", "unclear", "known",
             "duplicate_in_session"}
APPEND_ONLY = [
    "AGENTS.md", "notes/literature/prior_boundary.md", "notes/literature/search_log.md",
    "notes/literature/audit_sessions.md", "notes/literature/literature_audit_part2.md",
    "notes/literature/model_fidelity_methods.md", "notes/literature/already_implied.md",
    "notes/phases/PHASE_4_methodology_threats.md", "notes/00_evidence_inventory.md",
    "notes/00_exposure_register.md",
]
INSERT_ONLY = {  # file -> dòng cũ được phép thay (đúng một lần)
    "README.md": [], "notes/literature/README.md": [],
    "papers/README.md": ["| Cùng part, còn vướng truy cập | **P02 — Poelman2023** |"],
    "notes/literature/matrix.md": ["| P02 | Sensitivity theo prediction errors/horizon/update frequency"],
}
UNCHANGED = [
    "notes/literature/claims.md", "notes/literature/prior_predictions.md",
    "notes/01_definitions_v1_2.md", "notes/01_claim_boundary_v1_2.md",
    "notes/design/grid_v1_2.json", "notes/design/grid_v1_2.md", "notes/02_decision_log.md",
    "notes/lessons/L3.3_equivalence_queue.md", "notes/literature/terms.md",
]


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def git_show(path: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT)


def main() -> None:
    # 1. Ledger ↔ raw
    ledger = [json.loads(x) for x in (OUT / "calls.jsonl").read_text(encoding="utf-8").splitlines()]
    for r in ledger:
        raw = ROOT / r["raw_path"]
        assert raw.exists() and sha(raw) == r["sha256"] and raw.stat().st_size == r["bytes"], r["raw_path"]
        assert not raw.read_bytes().startswith(b"%PDF"), raw
        if r["http_status"] == 200:
            recs, _ = parse_records(r["provider"], raw.read_bytes())
            assert len(recs) == r["n_records"], r["call_id"]
    status = Counter(str(r["http_status"]) for r in ledger)
    ok = {r["call_id"]: r for r in ledger if r["http_status"] == 200}

    # 2. CSV sàng ↔ raw theo rank
    screened = [c for c, r in ok.items() if r["kind"] in ("keyword", "forward")]
    rows_all = []
    for cid in screened:
        r = ok[cid]
        recs, _ = parse_records(r["provider"], (ROOT / r["raw_path"]).read_bytes())
        rows = list(csv.DictReader((OUT / "screen" / f"{cid}.csv").open(encoding="utf-8")))
        dec = json.loads((OUT / "screen" / f"{cid}.decisions.json").read_text(encoding="utf-8"))
        assert len(rows) == len(recs) == len(dec), cid
        for i, (row, x) in enumerate(zip(rows, recs), start=1):
            assert int(row["rank"]) == i and row["raw_sha256"] == r["sha256"]
            rid = str(x.get("id") or "").rsplit("/", 1)[-1] if r["provider"].startswith("openalex") else x.get("id")
            assert row["record_id"] == rid and row["title"] == (x.get("title") or ""), (cid, i)
            assert row["decision"] in DECISIONS and [row["decision"], row["tier"], row["reason"]] == dec[str(i)]
        rows_all += rows
    assert not (set(p.stem for p in (OUT / "screen").glob("*.csv")) - set(screened))
    by_region = Counter(ok[r["call_id"]]["region"] for r in rows_all)
    decisions = Counter(r["decision"] for r in rows_all)
    cands = [r for r in rows_all if r["decision"].startswith("candidate")]
    known = json.loads((OUT / "known_ids.json").read_text(encoding="utf-8"))
    known_set = {x.lower() for k in ("dois", "previously_screened_dois") for x in known[k]}
    known_set |= set(known["arxiv"]) | set(known["previously_screened_arxiv"])
    for c in cands:
        ident = (c["doi"] or "").lower()
        arx = re.sub(r"^10\.48550/arxiv\.", "", ident)
        assert ident not in known_set and arx not in known_set, c
    assert len(rows_all) == 157 and by_region == {"A1": 60, "A2": 57, "A3": 40}, (len(rows_all), by_region)
    assert decisions["candidate_T1"] == 0 and decisions["candidate_T2"] == 4

    # 3. P02 provenance/record
    acc = json.loads((OUT / "source_access.json").read_text(encoding="utf-8"))["sources"][0]
    pdf = Path(acc["pdf_path"])
    assert pdf.read_bytes().startswith(b"%PDF-") and sha(pdf) == acc["pdf_sha256"]
    assert not pdf.resolve().is_relative_to(ROOT)
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    assert int(re.search(r"^Pages:\s+(\d+)", info, re.M)[1]) == acc["pages"] == 20
    assert "10.1016/j.trc.2023.104174" in info
    assert sha(Path(acc["text_path"])) == acc["text_sha256"]
    note = (ROOT / "notes/literature/papers/P02.md").read_text(encoding="utf-8")
    assert "AI-T" in note and "chưa** L-self" in note and "§3.3.1 tr15" in note
    snap_note = OUT / "snapshots" / "P02_note_before_2026-10-26.md"
    assert snap_note.read_bytes() == git_show("notes/literature/papers/P02.md")

    # 4. Catalog + workbook: chỉ P02 đổi
    cat = json.loads((ROOT / "papers/data/catalog.json").read_text(encoding="utf-8"))
    old = json.loads((OUT / "snapshots/catalog_before_2026-10-26.json").read_text(encoding="utf-8"))
    assert (OUT / "snapshots/catalog_before_2026-10-26.json").read_bytes() == git_show("papers/data/catalog.json")
    ids = [p["id"] for p in cat["papers"]]
    assert ids == [p["id"] for p in old["papers"]] and len(ids) == 19
    changed_cat = [n["id"] for n, o in zip(cat["papers"], old["papers"]) if n != o]
    assert changed_cat == ["P02"]
    p02 = cat["papers"][ids.index("P02")]
    assert p02["author_reading"] == "Chưa có xác nhận" and p02["local_sha256"] == acc["pdf_sha256"]
    xsnap = OUT / "snapshots/paper_comparison_before_2026-10-26.xlsx"
    assert xsnap.read_bytes() == git_show("papers/paper_comparison.xlsx")
    ob, nb = load_workbook(xsnap), load_workbook(ROOT / "papers/paper_comparison.xlsx")
    rows = lambda b, s: {r[0]: list(r) for r in b[s].iter_rows(min_row=2, values_only=True)}
    ln_o, ln_n = rows(ob, "Danh_sach"), rows(nb, "Danh_sach")
    nt_o, nt_n = rows(ob, "Ghi_chu_doc"), rows(nb, "Ghi_chu_doc")
    assert all(ln_n[k][11] == v[11] for k, v in ln_o.items())  # mức tác giả đọc
    assert [k for k, v in nt_o.items() if nt_n[k] != v] == ["P02"]
    assert [k for k, v in ln_o.items() if ln_n[k] != v] == ["P02"]
    assert nt_n["P02"][13] == "AI (chưa kiểm)" and nt_n["P02"][8] == "Mục liên quan"
    guide_n = [r[0] for r in nb["Huong_dan"].iter_rows(min_row=2, values_only=True)]
    guide_o = [r[0] for r in ob["Huong_dan"].iter_rows(min_row=2, values_only=True)]
    assert guide_n == guide_o[:1] + ["L2.8 follow-up"] + guide_o[1:]
    ob.close(); nb.close()

    # 5. Lịch sử: append-only / insert-only / unchanged
    for f in APPEND_ONLY:
        assert (ROOT / f).read_bytes().startswith(git_show(f)), f
    for f, allowed in INSERT_ONLY.items():
        a = git_show(f).decode().splitlines(); b = (ROOT / f).read_text(encoding="utf-8").splitlines()
        reps = 0
        for tag, i1, i2, _, _ in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
            if tag == "equal" or tag == "insert":
                continue
            assert tag == "replace" and i2 - i1 == 1 and any(a[i1].startswith(x) for x in allowed), (f, tag, a[i1:i2])
            reps += 1
        assert reps == len(allowed), (f, reps)
    for f in UNCHANGED:
        assert (ROOT / f).read_bytes() == git_show(f), f
    hist = subprocess.check_output(["git", "ls-tree", "-r", "--name-only", BASE,
                                    "results/lessons/l2_8/part1", "results/lessons/l2_8/part2",
                                    "results/lessons/l2_7"], cwd=ROOT, text=True).split()
    for f in hist:
        assert (ROOT / f).read_bytes() == git_show(f), f

    # 6. Không PDF trong Git; link nội bộ
    assert not subprocess.check_output(["git", "ls-files", "*.pdf", "*.PDF"], cwd=ROOT).strip()
    links = 0
    docs = [ROOT / "README.md", ROOT / "AGENTS.md", ROOT / "papers/README.md",
            *OUT.glob("*.md"), *(ROOT / "notes/literature").rglob("*.md"),
            ROOT / "notes/phases/PHASE_4_methodology_threats.md"]
    for doc in docs:
        for link in re.findall(r"\]\(([^)]+)\)", doc.read_text(encoding="utf-8")):
            if "://" in link or link.startswith(("#", "mailto:")):
                continue
            target = unquote(link.split("#", 1)[0].strip("<>"))
            if (doc.parent / target).resolve() == (OUT / "verification.json").resolve():
                links += 1  # output của chính audit này, ghi ở cuối
                continue
            assert (doc.parent / target).exists(), (str(doc), link)
            links += 1

    result = dict(
        artifact_status="PASS", responsible="Doan Van Tai", date="2026-10-26", base_commit=BASE,
        ledger_attempts=len(ledger), ledger_calls=len({r["call_id"] for r in ledger}),
        http_status_counts=dict(status), screened_calls=len(screened),
        screened_records=len(rows_all), screened_by_region=dict(by_region),
        decisions=dict(decisions), new_candidates_T2=[c["doi"] or c["arxiv"] for c in cands],
        rank_level_raw_retained=True, stopping_rule_met=False, saturation_claimed=False,
        p02_read_level="AI-T", p02_pdf_sha256=acc["pdf_sha256"], p02_pages=20,
        catalog_papers=19, catalog_rows_changed=changed_cat, workbook_note_rows_changed=["P02"],
        author_reading_levels_changed=0, historical_files_checked=len(hist),
        append_only_files=len(APPEND_ONLY), local_links_checked=links, pdfs_in_git=0,
        old_audit_l2_8_part2="FAIL_BY_DESIGN_at_P02_row_assertion",
        grid_run=False, heldout_outcomes_used=False, L2_9="NOT_STARTED_BY_REQUEST",
        human_L_self="OPEN", PC_labels_changed=False, contract_changed=False,
    )
    (OUT / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()

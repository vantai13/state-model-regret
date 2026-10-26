"""Audit ghi nhận Gate L GO có điều kiện (26/10/2026).

Người phụ trách: Doan Van Tai. Kiểm: xác nhận nguyên văn có; record L-self ghi author-attested và
không bịa lời giải; Excel chỉ đổi đúng 6 mức tác giả; DL-043 đúng luật (thiếu L8 → GO có điều kiện);
closeout giữ hai mục tác giả; chỉ append/insert; link sống; không PDF trong Git.
"""
from __future__ import annotations
import difflib, json, re, subprocess
from pathlib import Path
from urllib.parse import unquote
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results/lessons/l2_9/gate_l"
BASE = "6a789e7a3e74c8c7ab99f5c48dbf1dbcdbb712f7"
IDS = {"WM26", "R02", "R04", "P02", "S11", "P12"}
APPEND_ONLY = ["notes/literature/prior_boundary.md", "notes/02_decision_log.md", "notes/01_claim_boundary_v1_2.md",
               "notes/00_evidence_inventory.md", "notes/00_exposure_register.md", "AGENTS.md"]
INSERT_ONLY = ["README.md", "notes/literature/README.md"]


def git_show(p: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{BASE}:{p}"], cwd=ROOT)


def main() -> None:
    att = json.loads((OUT / "author_attestation.json").read_text(encoding="utf-8"))
    assert att["answers_verbatim"]["Về nhãn của PC1–PC6 và Main RQ?"] == "Xác nhận nhãn tạm"
    assert att["L8_reader"] == "Chưa có"
    rec = (ROOT / "notes/literature/author_reading_record.md").read_text(encoding="utf-8")
    rows = [r for r in rec.splitlines() if re.match(r"^\| [1-7] \|", r)]
    assert len(rows) == 7 and all("author-attested" in r and "Ở vở của tác giả, chưa chép vào repo" in r for r in rows)
    cat = json.loads((ROOT / "papers/data/catalog.json").read_text(encoding="utf-8"))["papers"]
    old = json.loads((OUT / "catalog_before.json").read_text(encoding="utf-8"))["papers"]
    changed = {n["id"] for n, o in zip(cat, old) if n != o}
    assert changed == IDS and all(n["author_reading"] == "Mục liên quan" for n in cat if n["id"] in IDS)
    assert all({k: v for k, v in n.items() if k != "author_reading"} == {k: v for k, v in o.items() if k != "author_reading"}
               for n, o in zip(cat, old))
    ob, nb = load_workbook(OUT / "paper_comparison_before.xlsx"), load_workbook(ROOT / "papers/paper_comparison.xlsx")
    lv = lambda b: {r[0]: r[11] for r in b["Danh_sach"].iter_rows(min_row=2, values_only=True)}
    assert {k for k, v in lv(ob).items() if lv(nb)[k] != v} == IDS
    nt = lambda b: {r[0]: list(r) for r in b["Ghi_chu_doc"].iter_rows(min_row=2, values_only=True)}
    assert nt(ob) == nt(nb)
    dl = (ROOT / "notes/02_decision_log.md").read_text(encoding="utf-8")
    assert dl.count("### DL-043 —") == 1
    d43 = dl.split("### DL-043 —", 1)[1]
    for q in ["Gate L: GO có điều kiện", "**L8 ❌**", "author-attested", "A01 chỉ sau khi L3.3 đã commit", "trước A04", "HOLD Phase 4"]:
        assert q in d43, q
    co = (ROOT / "notes/phases/PHASE_2_closeout.md").read_text(encoding="utf-8")
    assert co.startswith("# Phase 2 Closeout — GO có điều kiện") and co.count("⟨Tác giả tự viết — chưa có.⟩") == 2
    pb = (ROOT / "notes/literature/prior_boundary.md").read_text(encoding="utf-8")
    sec = pb.split("## Gate L — nhãn được tác giả xác nhận (26/10/2026)", 1)[1]
    assert len([r for r in sec.splitlines() if re.match(r"^\| (PC[1-6]|Main RQ) \|", r)]) == 7
    for f in APPEND_ONLY:
        assert (ROOT / f).read_bytes().startswith(git_show(f)), f
    for f in INSERT_ONLY:
        a = git_show(f).decode().splitlines(); b = (ROOT / f).read_text(encoding="utf-8").splitlines()
        assert all(t in ("equal", "insert") for t, *_ in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes()), f
    assert not subprocess.check_output(["git", "ls-files", "*.pdf", "*.PDF"], cwd=ROOT).strip()
    links = 0
    docs = [ROOT / "README.md", ROOT / "AGENTS.md", *(d for d in OUT.glob("*.md") if not d.name.endswith("_before.md")), *(ROOT / "notes").rglob("*.md")]
    for doc in docs:  # *_before.md là snapshot nguyên văn của file ở nơi khác
        for link in re.findall(r"\]\(([^)]+)\)", doc.read_text(encoding="utf-8")):
            if "://" in link or link.startswith(("#", "mailto:")):
                continue
            t = unquote(link.split("#", 1)[0].strip("<>"))
            if (doc.parent / t).resolve() == (OUT / "verification.json").resolve():
                links += 1; continue
            assert (doc.parent / t).exists(), (str(doc), link); links += 1
    res = dict(artifact_status="PASS", responsible="Doan Van Tai", date="2026-10-26", base_commit=BASE,
               gate_L="GO có điều kiện", decision="DL-043", author_attested=["WM26", "R02", "R04", "P11", "P02", "S11", "P12", "T03", "labels"],
               excel_author_levels_set=sorted(IDS), L8="OPEN", closeout_author_sections="OPEN",
               phase3="OPEN: L3.3 before A01; GVHD memo before A04; Phase 4 HOLD (P0-V10)",
               local_links_checked=links, pdfs_in_git=0)
    (OUT / "verification.json").write_text(json.dumps(res, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(res, ensure_ascii=False))


if __name__ == "__main__":
    main()

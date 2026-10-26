"""Audit L2.6 CSVs and preservation against the session baseline.

Responsible: Doan Van Tai. Only writes part_a/verification.json. Requires the
local workbook backup and named PDFs recorded in source_access.json.
"""
import ast
import csv
from fractions import Fraction as F
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"results/lessons/l2_6/part_a"


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def rows(name):
    with (OUT/name).open() as f:return list(csv.DictReader(f))


def main():
    ages=rows("periodic_age.csv");assert len(ages)==9
    for r in ages:
        t,d=F(r["period_seconds"]),F(r["delay_seconds"])
        assert F(r["continuous_mean_age"])==d+t/2
        assert F(r["slot_start_mean_age"])==sum(d+i for i in range(int(t)))/t
    # Sum Poisson probabilities for an even number of switches: a different
    # derivation from the generator's matrix-power/formula comparison.
    mode=rows("mode_and_input_variance.csv");assert len(mode)==3
    max_parity_error=0.
    tree=ast.parse((ROOT/"src/smr/design.py").read_text())
    fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=="episode_law")
    node=next(n.value for n in fn.body if isinstance(n,ast.Assign) and any(isinstance(t,ast.Name) and t.id=="keep" for t in n.targets))
    keep_expr=compile(ast.Expression(body=node),"repo keep expression","eval")
    for r in mode:
        k=float(r["kappa"])
        for t,col in [(1.,"same_after_1s"),(8.,"same_after_8s")]:
            lam=t/k;prob=math.exp(-lam);even=prob
            for n in range(1,240):
                prob*=lam/n
                if n%2==0:even+=prob
            error=abs(even-float(r[col]));max_parity_error=max(max_parity_error,error)
            assert error<1e-12
        repo_keep=eval(keep_expr,{"math":math,"float":float,"dt":F(1,120),"cell":{"kappa":F(r["kappa"])}})
        assert abs(repo_keep-(1+math.exp(-2/120/k))/2)<1e-15
        assert abs(float(r["no_switch_during_1s"])-math.exp(-1/k))<1e-15
        assert abs(float(r["sd_packets"])**2-float(r["variance_packets_squared"]))<1e-12
    cases=rows("binary_identity.csv");assert len(cases)==125
    for r in cases:
        a,b,p=F(r["gap1"]),F(r["gap2"]),F(r["p1"])
        before=min(F(0),p*a+(1-p)*b)
        after=p*min(F(0),a)+(1-p)*min(F(0),b)
        assert F(r["value"])==before-after==F(r["L01"])-F(r["L11"])
        assert F(r["value"])==min(F(r["Epositive"]),F(r["Enegative"]))
    exercise=rows("gap_exercise.csv");assert len(exercise)==2
    for r in exercise:
        a,b=F(r["gap_hi"]),F(r["gap_lo"])
        gain=min(F(0),(a+b)/2)-(min(a,F(0))+min(b,F(0)))/2
        assert gain==F(r["GSgivenM1"])==F(r["L01"])-F(r["L11"])
    assert F(exercise[1]["GSgivenM1"])==F(1,4)
    aoii=rows("aoi_aoii_example.csv");assert len(aoii)==6
    assert [int(r["AoII_linear_indicator"]) for r in aoii]==[0,0,1,2,3,0]

    backup=Path.home()/"papers/state-model-regret/archive/2026-10-26_before_l2_6_part_a/paper_comparison.xlsx"
    old=load_workbook(backup);new=load_workbook(ROOT/"papers/paper_comparison.xlsx")
    records=lambda w:{r[0]:list(r) for r in w["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    before,after=records(old),records(new)
    assert len(before)==8 and set(after)-set(before)=={"S11"}
    assert all(after[k]==v for k,v in before.items())
    reading=lambda w:{r[0]:r[11] for r in w["Danh_sach"].iter_rows(min_row=2,values_only=True)}
    br,ar=reading(old),reading(new)
    assert all(ar[k]==v for k,v in br.items())
    old.close();new.close()
    protected=json.loads((OUT/"protected_before.json").read_text())
    assert all(sha(ROOT/p)==h for p,h in protected["files"].items())
    prior="notes/literature/prior_boundary.md"
    original=subprocess.check_output(["git","show",protected["baseline_commit"]+":"+prior])
    assert (ROOT/prior).read_bytes().startswith(original)
    source=json.loads((OUT/"source_access.json").read_text())
    pdf=Path(source["S11"]["path"])
    assert pdf.read_bytes().startswith(b"%PDF-") and sha(pdf)==source["S11"]["sha256"]
    assert not pdf.resolve().is_relative_to(ROOT)
    assert sha(Path(source["R32"]["existing_PDF"]))==source["R32"]["existing_PDF_sha256"]
    lookups=json.loads((OUT/"lookup_records.json").read_text())
    assert len(lookups["keyword_first"]["screened"])==10 and len(lookups["keyword_fallback"]["screened"])==18
    assert lookups["crossref_s11_correct"]["doi"].lower()!="10.1109/twc.2025.3593377"
    links=0
    docs=[ROOT/"README.md",ROOT/"AGENTS.md",ROOT/"papers/README.md",OUT/"report.md",*list((ROOT/"notes/literature").rglob("*.md"))]
    for p in docs:
        for link in re.findall(r"\]\(([^)]+)\)",p.read_text()):
            if "://" in link or link.startswith(("#","mailto:")):continue
            rel=unquote(link.split("#")[0].split(' "')[0])
            if not rel:continue
            dest=(p.parent/rel).resolve()
            if dest!=OUT/"verification.json":assert dest.exists(),(p,rel)
            links+=1
    sums=(ROOT/"results/lessons/part_a/SHA256SUMS").read_text().splitlines()
    assert len(sums)==52
    for line in sums:
        h,p=line.split("  ",1);assert sha(ROOT/p)==h,p
    result=json.loads((OUT/"toy_results.json").read_text())
    assert result["source_sha256"]==sha(ROOT/"experiments/lessons/l2_6_freshness.py")
    verification=dict(status="PASS",responsible="Doan Van Tai",generator_checks=result["check_count"],
                      periodic_cases=len(ages),binary_cases=len(cases),mode_cases=len(mode),gap_exercises=len(exercise),
                      max_independent_Poisson_parity_error=max_parity_error,repo_keep_expression_checked=True,
                      previous_workbook_records_preserved=len(before),catalog_rows=len(after),
                      protected_files_unchanged=len(protected["files"]),prior_boundary_old_prefix_unchanged=True,
                      local_links_checked=links,part_a_checksums_verified=len(sums),keyword_title_records=28,
                      formal_author_gate="pending",part_b_P13="not performed",grid_outcomes_computed=0)
    (OUT/"verification.json").write_text(json.dumps(verification,ensure_ascii=False,indent=2)+"\n")
    assert (OUT/"verification.json").exists()
    print(json.dumps(verification,ensure_ascii=False))


if __name__=="__main__":main()

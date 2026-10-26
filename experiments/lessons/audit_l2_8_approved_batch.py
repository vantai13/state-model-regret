"""Audit đợt đọc được duyệt 26/10/2026 (IA26, DTS25, CEN26, CEH23).

Người phụ trách: Doan Van Tai. Kiểm hash/danh tính PDF, trích dẫn trong note có trong text,
catalog/workbook chỉ thêm 4 hàng đầu và giữ 19 record + mức tác giả, nhãn tạm không đổi,
record L-self không bị chạm, các file lịch sử chỉ append, link sống, không PDF trong Git.
"""

from __future__ import annotations

import difflib
import hashlib
import json
import re
import subprocess
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results/lessons/l2_8/approved_batch"
BASE = "cd7ffacb05a60600c285fbb16d2a61df3463c76d"
CUR = Path.home() / "papers/state-model-regret/current"
NEW = ["IA26", "DTS25", "CEN26", "CEH23"]
IDENT = {"IA26": "arXiv:2609.06351v1", "DTS25": "10.1145/3769700.3771695",
         "CEN26": "arXiv:2608.01725v1", "CEH23": "arXiv:2308.13166v1"}
QUOTES = {"IA26": ["controlled queueing, and decentralized Markov decision pro- cesses lie beyond this baseline",
                   "Fresh-local information strictly outperforms stale-global information if and only if",
                   "additional information cannot hurt when it is available with the same freshness"],
          "DTS25": ["Planned Validation", "While the description remains conceptual"],
          "CEN26": ["saturates near last-value baselines", "for which 1.0 matches the naive last-value forecast",
                    "cuts the violation rate from 20.5% to 9.0%"],
          "CEH23": ["more re-solving is not always beneficial", "Cooper (2002), Secomandi (2008)",
                    "converges to 1 exponentially fast", "ϕ : Rnx × Rnw × Rnu → Rnx is an affine function"]}
APPEND_ONLY = ["notes/literature/prior_boundary.md", "notes/01_claim_boundary_v1_2.md", "notes/00_evidence_inventory.md",
               "notes/00_exposure_register.md", "notes/literature/search_log.md", "notes/literature/h1_candidates.md",
               "AGENTS.md"]
INSERT_ONLY = ["README.md", "notes/literature/README.md", "papers/README.md"]
UNCHANGED = ["notes/literature/author_reading_record.md", "notes/literature/claims.md",
             "notes/literature/prior_predictions.md", "notes/02_decision_log.md", "notes/01_definitions_v1_2.md"]


def git_show(path: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT)


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s)


def main() -> None:
    acc = json.loads((OUT / "source_access.json").read_text(encoding="utf-8"))
    intake = Path(acc["intake_dir"])
    assert not intake.resolve().is_relative_to(ROOT)
    for s in acc["sources"]:
        p = intake / s["file"]
        assert hashlib.sha256(p.read_bytes()).hexdigest() == s["sha256"], s["file"]
        info = subprocess.check_output(["pdfinfo", str(p)], text=True)
        assert int(re.search(r"^Pages:\s+(\d+)", info, re.M)[1]) == s["pages"], s["file"]
    for i in NEW:
        assert (CUR / f"{i}.pdf").resolve().parent == intake.resolve()
        text = norm((CUR / f"{i}.txt").read_text(encoding="utf-8"))
        assert IDENT[i] in text.replace("doi.org/", ""), i
        for q in QUOTES[i]:
            assert norm(q) in text, (i, q)
        note = (ROOT / f"notes/literature/papers/{i}.md").read_text(encoding="utf-8")
        assert "chưa L-self" in note and ("AI-T" in note or "AI-F" in note)

    cat = json.loads((ROOT / "papers/data/catalog.json").read_text(encoding="utf-8"))["papers"]
    old = json.loads((OUT / "catalog_before.json").read_text(encoding="utf-8"))["papers"]
    assert (OUT / "catalog_before.json").read_bytes() == git_show("papers/data/catalog.json")
    assert [p["id"] for p in cat[:4]] == NEW and cat[4:] == old and len(cat) == 23
    assert all(p["author_reading"] == "Chưa có xác nhận" for p in cat)
    assert (OUT / "paper_comparison_before.xlsx").read_bytes() == git_show("papers/paper_comparison.xlsx")
    ob, nb = load_workbook(OUT / "paper_comparison_before.xlsx"), load_workbook(ROOT / "papers/paper_comparison.xlsx")
    rows = lambda b, s: {r[0]: list(r) for r in b[s].iter_rows(min_row=2, values_only=True)}
    assert all(rows(nb, "Ghi_chu_doc")[k] == v for k, v in rows(ob, "Ghi_chu_doc").items())
    levels = [r[11] for r in nb["Danh_sach"].iter_rows(min_row=2, values_only=True)]
    assert len(levels) == 23 and set(levels) == {"Chưa có xác nhận"}
    ob.close(); nb.close()

    appr = json.loads((OUT / "author_approval.json").read_text(encoding="utf-8"))
    assert appr["author_words"].startswith("các paper này chưa được duyệt") and len(appr["approved_targeted_reading"]) == 4
    for f in APPEND_ONLY:
        assert (ROOT / f).read_bytes().startswith(git_show(f)), f
    for f in INSERT_ONLY:
        a = git_show(f).decode().splitlines(); b = (ROOT / f).read_text(encoding="utf-8").splitlines()
        assert all(t in ("equal", "insert") for t, *_ in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes()), f
    for f in UNCHANGED:
        assert (ROOT / f).read_bytes() == git_show(f), f
    pb = (ROOT / "notes/literature/prior_boundary.md").read_text(encoding="utf-8")
    sec = pb.split("## Đợt đọc được duyệt (IA26, DTS25, CEN26, CEH23)", 1)[1]
    kept = re.findall(r"\*\*(NARROW|KEEP câu chữ|KEEP) giữ\*\*", sec)
    assert len(kept) == 5, kept

    assert not subprocess.check_output(["git", "ls-files", "*.pdf", "*.PDF"], cwd=ROOT).strip()
    links = 0
    for doc in [ROOT / "README.md", ROOT / "AGENTS.md", ROOT / "papers/README.md", *OUT.glob("*.md"),
                *(ROOT / "notes").rglob("*.md")]:
        for link in re.findall(r"\]\(([^)]+)\)", doc.read_text(encoding="utf-8")):
            if "://" in link or link.startswith(("#", "mailto:")):
                continue
            target = unquote(link.split("#", 1)[0].strip("<>"))
            if (doc.parent / target).resolve() == (OUT / "verification.json").resolve():
                links += 1
                continue
            assert (doc.parent / target).exists(), (str(doc), link)
            links += 1
    result = dict(artifact_status="PASS", responsible="Doan Van Tai", date="2026-10-26", base_commit=BASE,
                  pdfs_hashed=len(acc["sources"]), targeted_read=NEW, quotes_verified=sum(map(len, QUOTES.values())),
                  catalog_papers=len(cat), old_records_unchanged=19, author_levels_unconfirmed=23,
                  provisional_labels_changed=0, author_L_self_records_added=0, local_links_checked=links,
                  pdfs_in_git=0, gate_L="HOLD")
    (OUT / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()

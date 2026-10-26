"""Audit artifact L2.9 Part 1 (26/10/2026): hồ sơ 8 cột, nhãn để trống.

Người phụ trách: Doan Van Tai. AI viết theo yêu cầu. Kiểm: cột Nhãn chưa bị điền,
claim chép nguyên văn claims.md, mẫu L-self chưa có ô tác giả nào được điền, trích
dẫn WM26/freshness có thật trong nguồn, chỉ append/insert so với base, PC/contract/DL
và workbook không đổi, link nội bộ sống, không PDF trong Git. Không chứng nhận nhãn,
L-self, Gate L hay novelty.
"""

from __future__ import annotations

import difflib
import hashlib
import json
import re
import subprocess
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results" / "lessons" / "l2_9" / "part1"
BASE = "86ec87801cef51fb76876ec942ee4a4bfee984c0"
LABEL = "⟨tác giả chấm⟩"
APPEND_ONLY = ["notes/literature/prior_boundary.md", "notes/literature/papers/WM26.md", "AGENTS.md",
               "notes/00_evidence_inventory.md", "notes/00_exposure_register.md"]
INSERT_ONLY = ["README.md", "notes/literature/README.md"]
UNCHANGED = ["notes/literature/claims.md", "notes/literature/prior_predictions.md",
             "notes/01_definitions_v1_2.md", "notes/01_claim_boundary_v1_2.md",
             "notes/02_decision_log.md", "notes/literature/terms.md", "notes/design/grid_v1_2.json",
             "papers/paper_comparison.xlsx", "papers/data/catalog.json",
             "notes/literature/search_log.md", "notes/literature/papers/P02.md"]
WM26_QUOTES = [
    "C imports one true pool-mean vector, whereas S imports the candidate-wise true centered responses",
    "zj uses candidate j's realized observations",
    "the true-future oracles are diagnostic tools rather than deployable corrections",
    "learn response corrections from available observations and actions, without oracle access at decision time",
    "Neither C/S gains nor MSE fractions are additive causal responsibilities",
    "with unequal true-future information",
]


def git_show(path: str) -> bytes:
    return subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT)


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("’", "'")).strip()


def cells(line: str) -> list[str]:
    line = line.strip().strip("|").replace("\\|", "\0")
    return [c.strip().replace("\0", "\\|") for c in line.split("|")]


def main() -> None:
    # 1. Append/insert/unchanged so với base
    for f in APPEND_ONLY:
        assert (ROOT / f).read_bytes().startswith(git_show(f)), f
    for f in INSERT_ONLY:
        a = git_show(f).decode().splitlines(); b = (ROOT / f).read_text(encoding="utf-8").splitlines()
        ops = difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes()
        assert all(t in ("equal", "insert") for t, *_ in ops), f
    for f in UNCHANGED:
        assert (ROOT / f).read_bytes() == git_show(f), f

    # 2. Hồ sơ 8 cột: 7 dòng, nhãn trống, claim nguyên văn
    pb = (ROOT / "notes/literature/prior_boundary.md").read_text(encoding="utf-8")
    sec = pb.split("## L2.9 Part 1 — hồ sơ 8 cột", 1)[1].split("### Trạng thái Gate L", 1)[0]
    rows = [cells(l) for l in sec.splitlines() if l.startswith("| **")]
    assert len(rows) == 7 and all(len(r) == 8 for r in rows), [len(r) for r in rows]
    assert all(r[6] == LABEL for r in rows)
    label_words = re.compile(r"\b(KEEP|NARROW|PIVOT|DROP|HOLD)\b")
    assert not any(label_words.search(r[6]) for r in rows)
    claims = {}
    for l in (ROOT / "notes/literature/claims.md").read_text(encoding="utf-8").splitlines():
        if re.match(r"^\| PC[1-6] ", l):
            c = cells(l); claims[c[0]] = c[1]
    for r in rows[:6]:
        pc = re.match(r"\*\*(PC[1-6])\*\* (.*)", r[0])
        assert pc and pc[2] == claims[pc[1]], (r[0], claims.get(pc[1] if pc else None))
    assert rows[6][0].startswith("**Main RQ**")

    # 3. Mẫu L-self: không ô tác giả nào được điền
    rec = (ROOT / "notes/literature/author_reading_record.md").read_text(encoding="utf-8")
    trows = [cells(l) for l in rec.splitlines() if re.match(r"^\| [1-7] \|", l)]
    assert len(trows) == 7 and all(len(r) == 10 for r in trows)
    assert all(all(c == "" for c in r[5:]) for r in trows), "author cells must stay empty"

    # 4. Workbook: 19 nguồn, mức tác giả chưa đổi
    book = load_workbook(ROOT / "papers/paper_comparison.xlsx")
    levels = [r[11] for r in book["Danh_sach"].iter_rows(min_row=2, values_only=True)]
    book.close()
    assert len(levels) == 19 and set(levels) == {"Chưa có xác nhận"}

    # 5. Trích dẫn có trong nguồn
    acc = json.loads((OUT / "source_access.json").read_text(encoding="utf-8"))["sources"][0]
    pdf, txt = Path(acc["pdf_path"]), Path(acc["text_path"])
    assert hashlib.sha256(pdf.read_bytes()).hexdigest() == acc["pdf_sha256"]
    assert hashlib.sha256(txt.read_bytes()).hexdigest() == acc["text_sha256"]
    wm_text = norm(txt.read_text(encoding="utf-8")).replace("- ", "-")
    wm_note = norm((ROOT / "notes/literature/papers/WM26.md").read_text(encoding="utf-8"))
    for q in WM26_QUOTES:
        assert norm(q) in wm_text, q
    assert sum(norm(q) in wm_note for q in WM26_QUOTES) >= 5
    fr = (ROOT / "notes/literature/freshness_decision_value.md").read_text(encoding="utf-8")
    assert "không phải F của D22" in fr
    snap = OUT / "WM26_note_before_2026-10-26.md"
    assert snap.read_bytes() == git_show("notes/literature/papers/WM26.md")

    # 6. Không PDF trong Git; link nội bộ
    assert not subprocess.check_output(["git", "ls-files", "*.pdf", "*.PDF"], cwd=ROOT).strip()
    links = 0
    docs = [ROOT / "README.md", ROOT / "AGENTS.md", ROOT / "papers/README.md", *OUT.glob("*.md"),
            *(ROOT / "notes/literature").rglob("*.md")]
    for doc in docs:
        for link in re.findall(r"\]\(([^)]+)\)", doc.read_text(encoding="utf-8")):
            if "://" in link or link.startswith(("#", "mailto:")):
                continue
            target = unquote(link.split("#", 1)[0].strip("<>"))
            if (doc.parent / target).resolve() == (OUT / "verification.json").resolve():
                links += 1
                continue
            assert (doc.parent / target).exists(), (str(doc), link)
            links += 1

    result = dict(artifact_status="PASS", responsible="Doan Van Tai", date="2026-10-26", base_commit=BASE,
                  dossier_rows=len(rows), label_cells_filled_by_AI=0, claims_verbatim=6,
                  author_record_rows=len(trows), author_cells_filled=0, author_levels_unconfirmed=19,
                  wm26_quotes_verified=len(WM26_QUOTES), wm26_current_info_only_branch_found=False,
                  local_links_checked=links, pdfs_in_git=0, gate_L="HOLD",
                  PC_labels_assigned=False, decision_log_changed=False, contract_changed=False,
                  human_L_self="OPEN", reviewer_view_stored_separately=True)
    (OUT / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()

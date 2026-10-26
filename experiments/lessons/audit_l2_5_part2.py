"""Read-only reconstruction of L2.5 Part2 CSVs and preservation checks.

Responsible: Doan Van Tai. Uses multinomial counts, independently of the
generator's dynamic program. Only writes part2/verification.json.
"""
from collections import defaultdict
import csv
from fractions import Fraction as F
import hashlib
from itertools import product
import json
from math import factorial
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT/"results/lessons/l2_5/part2"


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def csv_rows(name):
    with (OUT/name).open() as f:
        return list(csv.DictReader(f))


def main():
    costs = defaultdict(lambda: F(0))
    events = csv_rows("cost_events.csv")
    assert len(events) == 32
    for r in events:
        a,b = F(r["bA"]),F(r["bB"])
        if r["action"] == "St":
            qa,qb = max(F(0),a-1),max(F(0),b-3)
        else:
            qa,qb = max(F(0),a-3),max(F(0),b-1)
        assert (F(r["qA"]),F(r["qB"]),F(r["cost"])) == (qa,qb,qa+qb)
        assert F(r["probability"]) == F(1,4)
        costs[r["case"],r["mode"],r["action"]] += F(1,4)*(qa+qb)
    result = json.loads((OUT/"toy_results.json").read_text())
    losses = defaultdict(dict)
    for r in csv_rows("policies.csv"):
        case,cell = r["case"],r["cell"]
        def choice(vals): return "Sw" if vals["Sw"] < vals["St"] else "St"
        model = {m: {a:F(v) for a,v in result["cases"][case]["J0" if cell[1]=="0" else "J"][m].items()}
                 for m in ("hi","lo")}
        for m in model:
            forecast = model[m] if cell[0]=="1" else {a:sum(model[k][a] for k in model)/2 for a in ("St","Sw")}
            assert r[m] == choice(forecast)
        loss = (costs[case,"hi",r["hi"]]+costs[case,"lo",r["lo"]])/2
        assert loss == F(r["expected_cost"])
        losses[case][cell] = loss
    summary = csv_rows("summary.csv")
    for r in summary:
        l=losses[r["case"]]
        assert all(F(r["L"+k]) == v for k,v in l.items())
        expected={"GS":l["00"]-l["10"],"GM":l["00"]-l["01"],"GSM":l["00"]-l["11"],
                  "Gam":l["10"]+l["01"]-l["00"]-l["11"],"GSgM1":l["01"]-l["11"],"GMgS1":l["10"]-l["11"]}
        assert all(F(r[k])==v for k,v in expected.items())

    initial=(0,3,3,3)
    all_dist=defaultdict(dict)
    for r in csv_rows("server_distribution.csv"):
        q=tuple(int(r[f"q{i}"]) for i in range(1,5))
        assert q not in all_dist[r["policy"]]
        all_dist[r["policy"]][q]=F(r["probability"])
    for policy,p in [("random",[F(1,4)]*4),("stale_d2",[F(1,2)]+[F(1,6)]*3)]:
        expected={}
        for a,b,c in product(range(9),repeat=3):
            d=8-a-b-c
            if d<0: continue
            counts=(a,b,c,d)
            mass=F(factorial(8))
            for i,n in enumerate(counts): mass=mass*p[i]**n/factorial(n)
            expected[tuple(initial[i]+counts[i] for i in range(4))]=mass
        assert all_dist[policy]==expected
    assert all_dist["stale_shortest"]=={(8,3,3,3):F(1)}
    assert all_dist["fresh_shortest"]=={(5,4,4,4):F(1)}
    for r in csv_rows("server_summary.csv"):
        dist=all_dist[r["policy"]]
        assert sum(dist.values())==1
        assert all(F(r[f"Eq{i+1}"])==sum(q[i]*p for q,p in dist.items()) for i in range(4))
        assert F(r["E_max_q"])==sum(max(q)*p for q,p in dist.items())

    backup=Path.home()/"papers/state-model-regret/archive/2026-10-26_before_l2_5_part2/paper_comparison.xlsx"
    old=load_workbook(backup);new=load_workbook(ROOT/"papers/paper_comparison.xlsx")
    before={r[0]:list(r) for r in old["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    after={r[0]:list(r) for r in new["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    assert len(before)==7 and len(after)==8 and "P01" not in after
    assert all(after[k]==v for k,v in before.items())
    readings=lambda w:{r[0]:r[11] for r in w["Danh_sach"].iter_rows(min_row=2,values_only=True)}
    br,ar=readings(old),readings(new)
    assert all(ar[k]==v for k,v in br.items())
    old.close();new.close()
    protected=json.loads((OUT/"protected_before.json").read_text())
    assert all(sha(ROOT/p)==h for p,h in protected["files"].items())
    prior="notes/literature/prior_boundary.md"
    original=subprocess.check_output(["git","show",protected["baseline_commit"]+":"+prior])
    assert (ROOT/prior).read_bytes().startswith(original)
    sources=json.loads((OUT/"source_access.json").read_text())
    for name in ["M00","P01"]:
        p=Path(sources[name]["path"])
        assert p.is_file() and p.read_bytes().startswith(b"%PDF-") and sha(p)==sources[name]["sha256"]
        assert not p.resolve().is_relative_to(ROOT)
    assert sha(Path(sources["R04_existing_provenance"]["note_path"]))==sources["R04_existing_provenance"]["note_sha256"]
    docs=[ROOT/"README.md",*list((ROOT/"notes/literature").rglob("*.md")),ROOT/"papers/README.md",OUT/"report.md"]
    links=0
    for p in docs:
        for link in re.findall(r"\]\(([^)]+)\)",p.read_text()):
            if "://" in link or link.startswith(("#","mailto:")): continue
            rel=unquote(link.split("#")[0].split(" \"")[0])
            if not rel: continue
            target=(p.parent/rel).resolve()
            # This auditor creates its own report after verifying the inputs.
            if target != OUT/"verification.json":
                assert target.exists(),(p,rel)
            links+=1
    checksums=(ROOT/"results/lessons/part_a/SHA256SUMS").read_text().splitlines()
    assert len(checksums)==52
    for line in checksums:
        h,p=line.split("  ",1)
        assert sha(ROOT/p)==h,p
    verification=dict(status="PASS",responsible="Doan Van Tai",generator_checks=result["check_count"],
                      cost_event_rows=len(events),policy_rows=8,summary_rows=len(summary),
                      terminal_distribution_rows=sum(len(d) for d in all_dist.values()),
                      independent_distribution_method="multinomial closed form, compared to generator DP",
                      previous_workbook_records_preserved=len(before),catalog_rows=len(after),
                      protected_files_unchanged=len(protected["files"]),prior_boundary_old_prefix_unchanged=True,
                      local_markdown_links_checked=links,part_a_checksums_verified=len(checksums),
                      source_PDFs_outside_repo=2,formal_author_gate="pending")
    (OUT/"verification.json").write_text(json.dumps(verification,ensure_ascii=False,indent=2)+"\n")
    assert (OUT/"verification.json").exists()
    print(json.dumps(verification,ensure_ascii=False))


if __name__=="__main__":
    main()

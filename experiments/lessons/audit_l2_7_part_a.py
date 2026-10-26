"""Replay exported L2.7 tables; verify provenance and preservation.

Responsible: Doan Van Tai. The oracle audit enumerates all contingent policies,
independently of the generator's pointwise-minimum implementation.
"""
import csv
from fractions import Fraction as F
import hashlib
from itertools import product
import json
from pathlib import Path
import re
import subprocess
from urllib.parse import unquote

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results/lessons/l2_7/part_a"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(name):
    with (OUT/name).open() as f:
        return list(csv.DictReader(f))


def main():
    result = json.loads((OUT/"toy_results.json").read_text())
    config = json.loads((OUT/"config.json").read_text())
    assert result["status"] == "PASS" and result["check_count"] == len(result["checks"]) == 14
    assert config["seed"] is None and config["tie_rule"] == "Stay"
    assert result["source_sha256"] == sha(ROOT/"experiments/lessons/l2_7_model_methods.py")
    assert result["input"]["sha256"] == sha(ROOT/result["input"]["path"])
    oracle_rows = read_csv("oracle_sweep.csv")
    assert len(oracle_rows) == 256
    seen_oracles = set()
    for row in oracle_rows:
        costs = [[F(row[k]) for k in keys] for keys in
                 [("stay_w1","stay_w2"),("switch_w1","switch_w2")]]
        seen_oracles.add(tuple(v for pair in costs for v in pair))
        contingent = {(a,b): (costs[a][0]+costs[b][1])/2
                      for a,b in product(range(2), repeat=2)}
        rp = min(contingent[0,0],contingent[1,1])
        ws = min(contingent.values())
        assert (F(row["RP"]),F(row["WS"]),F(row["future_EVPI"])) == (rp,ws,rp-ws)
        assert row["action"] == ("Stay" if contingent[0,0] <= contingent[1,1] else "Switch")
    assert seen_oracles == set(product(range(4),repeat=4))
    factorial_rows = read_csv("factorial_sweep.csv")
    assert len(factorial_rows) == 256
    seen_tables = set()
    for row in factorial_rows:
        a,b,c,d = [F(row[k]) for k in ("L00","L10","L01","L11")]
        seen_tables.add((a,b,c,d))
        betaS,betaM,betaSM = b-a,c-a,d-b-c+a
        assert F(row["Gamma"]) == -betaSM
        assert F(row["GS"]) == -betaS and F(row["GM"]) == -betaM
        assert F(row["GSgivenM1"]) == -betaS-betaSM
        assert F(row["GMgivenS1"]) == -betaM-betaSM
        assert F(row["averaged_state_effect"]) == -betaS-betaSM/2
        assert F(row["averaged_model_effect"]) == -betaM-betaSM/2
        assert F(row["additive_prediction_error"]) == -betaSM
    assert seen_tables == set(product(range(4),repeat=4))
    examples = read_csv("oracle_examples.csv")
    assert len(examples) == 4 and len(read_csv("factorial_examples.csv")) == 4
    for row in examples:
        cost,rp,ws = [F(row[k]) for k in ("expected_cost","RP","WS")]
        assert F(row["avoidable_regret"]) == cost-rp
        assert F(row["hindsight_gap"]) == cost-ws
        assert F(row["future_EVPI"]) == rp-ws

    protected = json.loads((OUT/"protected_before.json").read_text())
    assert len(protected["files"]) == 167
    for name,digest in protected["files"].items():
        assert sha(ROOT/name) == digest, name
    for name in ["notes/literature/prior_boundary.md","notes/literature/matrix.md",
                 "notes/literature/already_implied.md","notes/literature/search_log.md",
                 "notes/literature/h1_candidates.md"]:
        old = subprocess.check_output(["git","show",protected["baseline_commit"]+":"+name])
        assert (ROOT/name).read_bytes().startswith(old), name

    backup = Path.home()/"papers/state-model-regret/archive/2026-10-26_before_l2_7_part_a"
    old = load_workbook(backup/"paper_comparison.xlsx")
    new = load_workbook(ROOT/"papers/paper_comparison.xlsx")
    notes = lambda w: {r[0]:list(r) for r in w["Ghi_chu_doc"].iter_rows(min_row=2,values_only=True)}
    progress = lambda w: {r[0]:r[11] for r in w["Danh_sach"].iter_rows(min_row=2,values_only=True)}
    before,after = notes(old),notes(new)
    br,ar = progress(old),progress(new)
    assert len(before)==10 and len(after)==12 and set(after)-set(before)=={"P02","P07"}
    assert all(ar[k]==v for k,v in br.items())
    assert all(after[k]==v for k,v in before.items() if k!="R51")
    migration = json.loads((OUT/"workbook_note_update.json").read_text())
    assert before["R51"]==migration["before"] and after["R51"]==migration["after"]
    assert before["R51"][13]==after["R51"][13]=="AI (chưa kiểm)"
    assert before["R51"][8]=="Thứ cấp" and after["R51"][8]=="Mục liên quan"
    assert after["P07"][8]=="Mục liên quan" and after["P02"][8]=="Abstract"
    old.close();new.close()
    oldcat = json.loads((backup/"catalog.json").read_text())["papers"]
    catalog = json.loads((ROOT/"papers/data/catalog.json").read_text())["papers"]
    byid = {r["id"]:r for r in catalog}
    assert len(byid)==12 and "P03" not in byid and "R51" in byid
    assert all(byid[r["id"]]==r for r in oldcat if r["id"]!="R51")
    assert byid["P02"]["local"]=="" and byid["P02"]["ai"]=="Abstract"

    access = json.loads((OUT/"source_access.json").read_text())
    assert access["P02"]["full_text_status"]=="BLOCKED" and access["P02"]["pdf_path"] is None
    assert len(access["sources"])==2
    renders = 0
    for s in access["sources"]:
        p = Path(s["pdf_path"])
        assert p.read_bytes().startswith(b"%PDF-") and sha(p)==s["sha256"]
        assert not p.resolve().is_relative_to(ROOT)
        for render in s["rendered_pages"]:
            assert sha(Path(render["path"]))==render["sha256"]
            renders += 1
    assert renders==4
    sums = (ROOT/"results/lessons/part_a/SHA256SUMS").read_text().splitlines()
    assert len(sums)==52
    for line in sums:
        digest,name = line.split("  ",1)
        assert sha(ROOT/name)==digest,name
    docs = [ROOT/"README.md",ROOT/"AGENTS.md",ROOT/"papers/README.md",OUT/"report.md",
            ROOT/"notes/phases/PHASE_4_methodology_threats.md",
            *list((ROOT/"notes/literature").rglob("*.md"))]
    links = 0
    for doc in docs:
        for link in re.findall(r"\]\(([^)]+)\)",doc.read_text()):
            if "://" in link or link.startswith(("#","mailto:")):
                continue
            path = unquote(link.split("#",1)[0].strip("<>"))
            assert (doc.parent/path).exists(),(doc,link)
            links += 1
    assert not subprocess.check_output(["git","ls-files","*.pdf"]).strip()
    verification = dict(artifact_status="PASS",lesson_status="PARTIAL_P02_FULL_TEXT_BLOCKED",
                        responsible="Doan Van Tai",generator_checks=14,
                        independent_csv_replay=dict(oracle=256,factorial=256,examples=8),
                        protected_files_unchanged=167,legacy_manifest_files=52,
                        old_author_progress_preserved=10,other_old_notes_preserved=9,
                        R51_AI_note_migration_verified=True,catalog_unique_papers=12,
                        primary_targeted_sources=2,abstract_only_blocked_sources=1,
                        rendered_pages_verified=renders,relative_links_checked=links,
                        grid_run=False,null_selected=False,human_gate="OPEN")
    (OUT/"verification.json").write_text(json.dumps(verification,ensure_ascii=False,indent=2)+"\n")
    print(json.dumps(verification,ensure_ascii=False))


if __name__ == "__main__":
    main()

"""Cập nhật record S11, R01, R34 sau khi tác giả tải bản chính thức/bản khác (26/10/2026).

Người phụ trách: Doan Van Tai. Chỉ sửa 3 ID; dừng nếu dòng ghi chú không do AI ghi hoặc
tác giả đã chọn mức đọc. Snapshot: results/lessons/l2_9/journal_check/.
"""
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
CAT, XLSX = ROOT / "papers/data/catalog.json", ROOT / "papers/paper_comparison.xlsx"
CUR = Path.home() / "papers/state-model-regret/current"
UPD = {
    "S11": dict(access="Có (journal TIT 72(4) 40 trang + arXiv v3)", local="S11_TIT2026.pdf",
                local_source="Tác giả tải bản IEEE 26/10/2026",
                role="Journal TIT2026 đã có; số mục §III/Eq10, §VII-D/Theorem9/Cor2/AppendixO khớp arXiv v3 (AI kiểm 26/10); chưa L-self.",
                next="Tác giả đọc §III tr5–6 (Problem1 Eq10) bản journal."),
    "R01": dict(access="Có TR2006 (local); bản eCommons tác giả tải 26/10 = cùng TR; journal Annals OR chưa có",
                role="TR2006 AI-T; số mục bản journal 2010 **chưa kiểm** (cần bản Springer).",
                next="Cần bản journal Annals OR 177:91–114 để đối chiếu Prop1/Example3."),
    "R34": dict(access="TAC 1995 chưa có; có bản CDC 1992 (2 trang) tác giả tải 26/10",
                role="Bản CDC 1992 đọc 2 trang: JSEQ tối ưu khi trễ k=1, không tối ưu khi k≥2; cần TAC 1995 để kiểm chứng/phản ví dụ.",
                result="CDC 1992 tr997–998: Thm2.1.2 JSEQ tối ưu với k=1; k≥2 có phản ví dụ (trỏ [4]); admission một queue: hai ngưỡng (Thm3.1.2).",
                limit="Chỉ bản hội nghị 2 trang, chứng minh lược; mô hình đúng (không trục model).",
                next="Cần bản TAC 40(8) 1995."),
}


def main():
    data = json.loads(CAT.read_text(encoding="utf-8"))
    book = load_workbook(XLSX)
    lst, notes = book["Danh_sach"], book["Ghi_chu_doc"]
    for p in data["papers"]:
        if p["id"] not in UPD:
            continue
        assert p["author_reading"] == "Chưa có xác nhận"
        u = UPD[p["id"]]
        for k, v in u.items():
            p[k] = v
        if "local" in u:
            p["local_sha256"] = hashlib.sha256((CUR / u["local"]).read_bytes()).hexdigest()
        row = p["note_row"]
        if "result" in u:
            row[4] = u["result"]
        if "limit" in u:
            row[5] = u["limit"]
        row[12], row[14] = u["next"], "2026-10-26"
        rl = next(r for r in range(2, lst.max_row + 1) if lst.cell(r, 1).value == p["id"])
        assert lst.cell(rl, 12).value == "Chưa có xác nhận"
        rn = next(r for r in range(2, notes.max_row + 1) if notes.cell(r, 1).value == p["id"])
        assert notes.cell(rn, 14).value == "AI (chưa kiểm)"
        for c, v in enumerate(row, start=1):
            notes.cell(rn, c).value = v
    CAT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    book.save(XLSX)
    book.close()
    print(json.dumps({"updated": sorted(UPD)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
